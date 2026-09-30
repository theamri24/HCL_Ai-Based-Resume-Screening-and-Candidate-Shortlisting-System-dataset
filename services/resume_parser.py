"""
services/resume_parser.py
-------------------------
Extract text from PDF (PyMuPDF) and DOCX (python-docx).
Also handles file validation, secure filename, saving uploads.
"""

import os
import re
import uuid
from pathlib import Path

# PDF
try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

# DOCX
try:
    import docx
except ImportError:
    docx = None

from werkzeug.utils import secure_filename


# ------------------------------------------------------------
# Config
# ------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent.parent
UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED_EXTENSIONS = {"pdf", "docx"}
MAX_FILE_SIZE_MB = 10


# ------------------------------------------------------------
# Validation
# ------------------------------------------------------------
def allowed_file(filename: str) -> bool:
    """Check if file extension is allowed."""
    if "." not in filename:
        return False
    ext = filename.rsplit(".", 1)[1].lower()
    return ext in ALLOWED_EXTENSIONS


def validate_file_size(file_bytes: bytes) -> None:
    """Raise ValueError if file too big."""
    size_mb = len(file_bytes) / (1024 * 1024)
    if size_mb > MAX_FILE_SIZE_MB:
        raise ValueError(
            f"File too large ({size_mb:.1f} MB). Max allowed: {MAX_FILE_SIZE_MB} MB"
        )


def get_file_extension(filename: str) -> str:
    return filename.rsplit(".", 1)[1].lower() if "." in filename else ""


# ------------------------------------------------------------
# Save uploaded file
# ------------------------------------------------------------
def save_upload(file_bytes: bytes, original_filename: str) -> Path:
    """
    Save file to uploads/ with a unique secure name.
    Returns saved file path.
    """
    if not allowed_file(original_filename):
        raise ValueError(
            f"Unsupported file type. Allowed: {', '.join(ALLOWED_EXTENSIONS)}"
        )

    validate_file_size(file_bytes)

    ext = get_file_extension(original_filename)
    # secure base name + unique id
    base = secure_filename(original_filename.rsplit(".", 1)[0]) or "resume"
    unique_name = f"{base}_{uuid.uuid4().hex[:8]}.{ext}"
    save_path = UPLOAD_DIR / unique_name

    with open(save_path, "wb") as f:
        f.write(file_bytes)

    return save_path


# ------------------------------------------------------------
# PDF extraction
# ------------------------------------------------------------
def extract_text_from_pdf(file_path: str | Path) -> str:
    """Extract text using PyMuPDF."""
    if fitz is None:
        raise RuntimeError("PyMuPDF not installed. Run: pip install pymupdf")

    file_path = str(file_path)
    text_parts = []

    try:
        with fitz.open(file_path) as doc:
            if doc.page_count == 0:
                raise ValueError("PDF has no pages.")
            for page in doc:
                page_text = page.get_text("text")
                if page_text:
                    text_parts.append(page_text)
    except Exception as e:
        raise RuntimeError(f"Failed to read PDF: {e}")

    text = "\n".join(text_parts).strip()
    if not text:
        raise ValueError(
            "No readable text found in PDF. It may be a scanned image PDF."
        )
    return text


# ------------------------------------------------------------
# DOCX extraction
# ------------------------------------------------------------
def extract_text_from_docx(file_path: str | Path) -> str:
    """Extract text using python-docx (paragraphs + tables)."""
    if docx is None:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")

    file_path = str(file_path)
    text_parts = []

    try:
        document = docx.Document(file_path)

        # Paragraphs
        for para in document.paragraphs:
            if para.text and para.text.strip():
                text_parts.append(para.text.strip())

        # Tables (some resumes use tables)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():
                        text_parts.append(cell.text.strip())
    except Exception as e:
        raise RuntimeError(f"Failed to read DOCX: {e}")

    text = "\n".join(text_parts).strip()
    if not text:
        raise ValueError("No readable text found in DOCX.")
    return text


# ------------------------------------------------------------
# Main dispatcher
# ------------------------------------------------------------
def extract_resume_text(file_path: str | Path) -> str:
    """
    Auto-detect file type and extract text.
    Raises ValueError/RuntimeError on failure.
    """
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = file_path.suffix.lower().lstrip(".")

    if ext == "pdf":
        raw = extract_text_from_pdf(file_path)
    elif ext == "docx":
        raw = extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: .{ext}")

    return clean_text(raw)


# ------------------------------------------------------------
# Text cleaning (basic)
# ------------------------------------------------------------
def clean_text(text: str) -> str:
    """Basic cleanup — remove weird chars, extra whitespace."""
    if not text:
        return ""

    # Normalize line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Remove multiple blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)

    # Replace tabs with space
    text = text.replace("\t", " ")

    # Remove zero-width chars
    text = re.sub(r"[\u200b-\u200f\u202a-\u202e]", "", text)

    # Collapse multiple spaces
    text = re.sub(r" {2,}", " ", text)

    return text.strip()


# ------------------------------------------------------------
# CLI test
# ------------------------------------------------------------
if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("Usage: python services/resume_parser.py <path_to_resume.pdf|docx>")
        print("\nRunning self-test with dummy text...")
        sample = "John Doe\nPython, Machine Learning, SQL\n2 years experience"
        print("\n--- Sample cleaned text ---")
        print(clean_text(sample))
        print("\n--- Validation tests ---")
        print("allowed_file('x.pdf') :", allowed_file("x.pdf"))
        print("allowed_file('x.docx'):", allowed_file("x.docx"))
        print("allowed_file('x.txt') :", allowed_file("x.txt"))
        sys.exit(0)

    path = sys.argv[1]
    try:
        text = extract_resume_text(path)
        print(f"✅ Extracted {len(text)} characters from {path}\n")
        print("--- First 1000 chars ---")
        print(text[:1000])
    except Exception as e:
        print(f"❌ Error: {e}")